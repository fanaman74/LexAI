import pytest

from app import convert


def test_supported_extensions():
    assert convert.SUPPORTED_EXTENSIONS == {
        ".pdf", ".docx", ".doc", ".msg", ".eml", ".xlsx", ".csv", ".txt", ".rtf"}


def test_unsupported_raises():
    with pytest.raises(convert.ConversionError):
        convert.convert_to_markdown("photo.jpg", b"...")


def test_txt():
    result = convert.convert_to_markdown("note.txt", "héllo legal".encode())
    assert "héllo legal" in result.full_text and result.converter_used == "text"


def test_csv():
    result = convert.convert_to_markdown("t.csv", b"name,amount\nAlpha,100\n")
    assert "Alpha" in result.full_text and result.converter_used == "markitdown"


def test_docx(tmp_path):
    import docx
    p = tmp_path / "brief.docx"
    d = docx.Document()
    d.add_paragraph("The defendant breached the agreement.")
    d.save(p)
    result = convert.convert_to_markdown("brief.docx", p.read_bytes())
    assert "defendant breached" in result.full_text and result.converter_used == "markitdown"


def test_xlsx(tmp_path):
    import openpyxl
    p = tmp_path / "fees.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["item", "fee"])
    wb.active.append(["filing", 350])
    wb.save(p)
    result = convert.convert_to_markdown("fees.xlsx", p.read_bytes())
    assert "filing" in result.full_text and result.converter_used == "markitdown"


def test_eml():
    raw = (b"From: counsel@firm.com\r\nTo: client@example.com\r\n"
           b"Subject: Settlement offer\r\nDate: Mon, 1 Jun 2026 10:00:00 +0000\r\n"
           b"Content-Type: text/plain\r\n\r\nWe propose USD 50,000.\r\n")
    result = convert.convert_to_markdown("offer.eml", raw)
    assert "Settlement offer" in result.full_text and "50,000" in result.full_text
    assert result.converter_used == "eml"


def test_msg(monkeypatch, tmp_path):
    class FakeMsg:
        sender = "a@b.com"
        to = "c@d.com"
        subject = "Hearing date"
        date = "2026-06-01"
        body = "Hearing moved to July 3."
        recipients = []
        attachments = []
        def close(self):
            pass
    monkeypatch.setattr(convert.extract_msg, "Message", lambda path: FakeMsg())
    result = convert.convert_to_markdown("mail.msg", b"fakebinary")
    assert "Hearing date" in result.full_text and "July 3" in result.full_text
    assert result.converter_used == "msg"


def test_rtf_via_textutil():
    raw = rb"{\rtf1\ansi This clause is governed by Swiss law.}"
    result = convert.convert_to_markdown("clause.rtf", raw)
    assert "Swiss law" in result.full_text and result.converter_used == "textutil"


import shutil
import subprocess


def _make_text_pdf(path, text):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.multi_cell(0, 8, text)
    pdf.output(str(path))


def _make_blank_pdf(path):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.output(str(path))


LONG_TEXT = ("This Services Agreement is entered into between Alpha Corp and "
             "Beta LLC on January 5, 2026, and is governed by the laws of the "
             "State of Delaware. Each party shall keep all information confidential.")


def test_pdf_with_text_layer(tmp_path):
    p = tmp_path / "contract.pdf"
    _make_text_pdf(p, LONG_TEXT)
    result = convert.convert_to_markdown("contract.pdf", p.read_bytes())
    assert "Alpha Corp" in " ".join(result.full_text.split())
    assert result.converter_used in ("pdfplumber", "markitdown")


def test_pdf_without_text_triggers_ocr(tmp_path, monkeypatch):
    blank = tmp_path / "scan.pdf"
    _make_blank_pdf(blank)
    ocred = tmp_path / "ocred.pdf"
    _make_text_pdf(ocred, LONG_TEXT)

    def fake_run(cmd, **kwargs):
        assert cmd[0] == "ocrmypdf"
        shutil.copy(ocred, cmd[-1])
        return subprocess.CompletedProcess(cmd, 0, stdout="", stderr="")

    monkeypatch.setattr(convert.subprocess, "run", fake_run)
    result = convert.convert_to_markdown("scan.pdf", blank.read_bytes())
    assert "Alpha Corp" in " ".join(result.full_text.split()) and result.converter_used == "ocr"


def test_ocr_failure_raises_conversion_error(tmp_path, monkeypatch):
    blank = tmp_path / "scan.pdf"
    _make_blank_pdf(blank)

    def fake_run(cmd, **kwargs):
        return subprocess.CompletedProcess(cmd, 1, stdout="", stderr="bad pdf")

    monkeypatch.setattr(convert.subprocess, "run", fake_run)
    with pytest.raises(convert.ConversionError, match="OCR failed"):
        convert.convert_to_markdown("scan.pdf", blank.read_bytes())
